import os
from pathlib import Path
from datetime import datetime

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from urllib.parse import quote_plus


load_dotenv()


def get_engine():
    server = os.getenv("DB_SERVER", "").strip()
    database = os.getenv("DB_NAME", "").strip()
    driver = os.getenv("DB_DRIVER", "ODBC Driver 17 for SQL Server").strip()
    trusted = os.getenv("DB_TRUSTED_CONNECTION", "no").strip().lower()

    if trusted in ["yes", "true", "1"]:
        conn = (
            f"DRIVER={{{driver}}};"
            f"SERVER={server};"
            f"DATABASE={database};"
            f"Trusted_Connection=yes;"
            f"TrustServerCertificate=yes;"
        )
    else:
        user = os.getenv("DB_USER", "").strip()
        password = os.getenv("DB_PASSWORD", "").strip()

        conn = (
            f"DRIVER={{{driver}}};"
            f"SERVER={server};"
            f"DATABASE={database};"
            f"UID={user};"
            f"PWD={password};"
            f"TrustServerCertificate=yes;"
        )

    return create_engine(
        "mssql+pyodbc:///?odbc_connect=" + quote_plus(conn),
        fast_executemany=True,
    )


def clean_text(value):
    return str(value or "").strip()


def detect_faq_category(question, excel_category=""):
    text = clean_text(question)

    accounting_keywords = [
        "حسابداری",
        "سند حسابداری",
        "ثبت سند",
        "سند",
        "تراز",
        "کل",
        "معین",
        "تفصیل",
        "تفصیلی",
        "درآمد",
        "هزینه",
        "سود و زیان",
        "اختتامیه",
        "افتتاحیه",
        "مودیان",
        "سامانه مودیان",
        "صورتحساب",
        "حافظه مالیاتی",
        "شناسه یکتا",
        "شماره اقتصادی",
        "گزارش اطلاعات فاقد سند",
    ]

    system_keywords = [
        "ورود",
        "رمز عبور",
        "شرکت",
        "سال مالی",
        "کاربر",
        "دسترسی",
        "نقش",
        "شیفت",
        "پیام",
        "مانیتور کاربران",
        "مراحل",
        "تأیید",
        "تایید",
        "لوگو",
        "تم",
        "مرورگر",
        "گروه‌بندی شرکت",
        "گروه بندی شرکت",
        "صفحه شخصی",
    ]

    for keyword in accounting_keywords:
        if keyword in text:
            return "حسابداری"

    for keyword in system_keywords:
        if keyword in text:
            return "مدیریت سیستم"

    category = clean_text(excel_category)
    return category or "عمومی"


def find_excel_file():
    current_dir = Path(__file__).resolve().parent

    candidates = list(current_dir.glob("*.xlsx"))

    if not candidates:
        raise FileNotFoundError("هیچ فایل اکسل xlsx کنار این فایل پیدا نشد.")

    for file in candidates:
        if "Medrik_System_Management_FAQ" in file.name:
            return file

    return candidates[0]


def get_columns(conn):
    rows = conn.execute(
        text(
            """
            SELECT COLUMN_NAME
            FROM INFORMATION_SCHEMA.COLUMNS
            WHERE TABLE_SCHEMA = 'dbo'
              AND TABLE_NAME = 'knowledge_base'
            """
        )
    ).fetchall()

    return {row[0] for row in rows}


def find_header_indexes(header_row):
    headers = {}

    for index, cell in enumerate(header_row):
        title = clean_text(cell.value)
        if title:
            headers[title] = index

    question_col = None
    answer_col = None
    category_col = None
    keywords_col = None

    for key, index in headers.items():
        if key in ["سؤال", "سوال", "question"]:
            question_col = index
        elif key in ["پاسخ", "answer"]:
            answer_col = index
        elif key in ["دسته", "category"]:
            category_col = index
        elif key in ["کلیدواژه‌ها", "کلیدواژه ها", "keywords"]:
            keywords_col = index

    if question_col is None or answer_col is None:
        raise ValueError("ستون سؤال و پاسخ در اکسل پیدا نشد.")

    return question_col, answer_col, category_col, keywords_col


def main():
    excel_path = find_excel_file()
    print("Excel file:", excel_path.name)

    wb = load_workbook(excel_path)
    ws = wb.active

    question_col, answer_col, category_col, keywords_col = find_header_indexes(ws[1])

    engine = get_engine()
    now = datetime.now().isoformat(timespec="seconds")

    inserted = 0
    updated = 0
    skipped = 0
    accounting_count = 0
    system_count = 0

    with engine.begin() as conn:
        columns = get_columns(conn)

        for row in ws.iter_rows(min_row=2):
            question = clean_text(row[question_col].value)
            answer = clean_text(row[answer_col].value)

            if not question or not answer:
                skipped += 1
                continue

            excel_category = ""
            if category_col is not None:
                excel_category = clean_text(row[category_col].value)

            keywords = ""
            if keywords_col is not None:
                keywords = clean_text(row[keywords_col].value)

            category = detect_faq_category(question, excel_category)

            if category == "حسابداری":
                accounting_count += 1
            elif category == "مدیریت سیستم":
                system_count += 1

            existing = conn.execute(
                text(
                    """
                    SELECT TOP 1 id
                    FROM dbo.knowledge_base
                    WHERE question = :question
                    """
                ),
                {"question": question},
            ).fetchone()

            if existing:
                set_parts = []
                params = {
                    "id": existing[0],
                    "question": question,
                    "answer": answer,
                    "category": category,
                    "keywords": keywords,
                    "is_active": 1,
                    "updated_at": now,
                }

                if "answer" in columns:
                    set_parts.append("answer = :answer")
                if "category" in columns:
                    set_parts.append("category = :category")
                if "keywords" in columns:
                    set_parts.append("keywords = :keywords")
                if "is_active" in columns:
                    set_parts.append("is_active = :is_active")
                if "updated_at" in columns:
                    set_parts.append("updated_at = :updated_at")

                if set_parts:
                    sql = f"""
                    UPDATE dbo.knowledge_base
                    SET {", ".join(set_parts)}
                    WHERE id = :id
                    """
                    conn.execute(text(sql), params)
                    updated += 1

            else:
                insert_cols = []
                values = []
                params = {
                    "question": question,
                    "answer": answer,
                    "category": category,
                    "keywords": keywords,
                    "is_active": 1,
                    "created_at": now,
                    "updated_at": now,
                }

                for col in ["question", "answer", "category", "keywords", "is_active", "created_at", "updated_at"]:
                    if col in columns:
                        insert_cols.append(col)
                        values.append(f":{col}")

                sql = f"""
                INSERT INTO dbo.knowledge_base ({", ".join(insert_cols)})
                VALUES ({", ".join(values)})
                """

                conn.execute(text(sql), params)
                inserted += 1

    print("--------------------------------")
    print("Import finished.")
    print("Inserted:", inserted)
    print("Updated:", updated)
    print("Skipped:", skipped)
    print("Accounting category:", accounting_count)
    print("System management category:", system_count)
    print("--------------------------------")


if __name__ == "__main__":
    main()